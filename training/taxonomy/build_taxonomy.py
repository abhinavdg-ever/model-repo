#!/usr/bin/env python3
"""Build taxonomy_v24.json from the client's Page_Taxonomy.xlsx.

    python build_taxonomy.py                      # Page_Taxonomy.xlsx beside this file
    python build_taxonomy.py ~/Desktop/Page_Taxonomy.xlsx
    python build_taxonomy.py --check              # fail if a copy is out of date

The workbook is the single source of the vocabulary: L1 > L2 > L3 with
descriptions, the specialty list, and a codeability tag per path. The JSON is
written here and copied into annotation-tool/ and bert-training/, which each
stand alone and read only their own copy. Neither holds a page type name in
code, so a new page type is an edit to the workbook and a re-run of this script.

Structural rules that the workbook states only in prose (which L1 takes its L2
from the content section, the names of the fallback values) are the options
below, and are written into the JSON's ``rules`` block for both readers.
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
TRAINING = HERE.parent
OUT_NAME = "taxonomy_v24.json"
COPIES = (HERE / OUT_NAME, TRAINING / "annotation-tool" / OUT_NAME, TRAINING / "bert-training" / OUT_NAME)

SHEET_PATHS = "Taxonomy L1-L2-L3"
SHEET_SPECIALTY = "Specialty"
SHEET_L1 = "L1 Description"

# Workbook codeability wording → the four codeability values.
CODEABILITY = {
    "codeable": "codeable",
    "non-codeable": "non_codeable",
    "discharge": "discharge",
    "needs review": "needs_review",
}


def codeability_key(raw: str) -> str:
    """'Context dependent (...)' and anything unrecognised is needs_review,
    never non_codeable: silently dropping a page is the expensive error."""
    text = (raw or "").strip().lower()
    return CODEABILITY.get(text, "needs_review")


def clean(value) -> str:
    return " ".join(str(value).split()) if value is not None else ""


def build(xlsx: Path, *, section_l1: str, fallback: str, unknown: str) -> dict:
    import openpyxl

    book = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    for sheet in (SHEET_PATHS, SHEET_SPECIALTY, SHEET_L1):
        if sheet not in book.sheetnames:
            raise SystemExit(f"{xlsx.name}: no sheet {sheet!r} (has {book.sheetnames})")

    l1_desc = {clean(r[0]): clean(r[1]) for r in book[SHEET_L1].iter_rows(min_row=2, values_only=True) if r[0]}
    specialty = [{"name": clean(r[0]), "description": clean(r[1])}
                 for r in book[SHEET_SPECIALTY].iter_rows(min_row=2, values_only=True) if r[0]]

    l1_order: list[str] = []
    l1_extra: dict[str, dict] = {}
    l2: dict[str, list[dict]] = {}
    l3: dict[str, list[dict]] = {}
    seen: set[tuple[str, str, str]] = set()
    for number, row in enumerate(book[SHEET_PATHS].iter_rows(min_row=2, values_only=True), 2):
        if not row or not row[0]:
            continue
        a, b, c = clean(row[0]), clean(row[1]), clean(row[2])
        key = (a, b, c)
        if key in seen:
            raise SystemExit(f"{SHEET_PATHS} row {number}: duplicate path {' > '.join(filter(None, key))}")
        seen.add(key)
        entry = {
            "description": clean(row[3]),
            "specialty": clean(row[4]),
            "codeability": codeability_key(clean(row[5])),
            "codeability_tag": clean(row[5]),
            "notes": clean(row[6]) if len(row) > 6 else "",
        }
        if a not in l1_order:
            l1_order.append(a)
            l2[a] = []
        if not b:
            if c:
                raise SystemExit(f"{SHEET_PATHS} row {number}: an L3 with no L2")
            l1_extra[a] = entry  # an L1 with no L2 (Problem List, Others)
        elif not c:
            l2[a].append({"name": b, **entry})
        else:
            if not any(x["name"] == b for x in l2[a]):
                # The L3 rows carry the L2's description after the L3 name.
                l2[a].append({"name": b, **{**entry, "description": entry["description"].split(". ", 1)[-1],
                                            "specialty": ""}})
            l3.setdefault(f"{a} > {b}", []).append({"name": c, **entry})

    missing = [name for name in l1_order if name not in l1_desc]
    if missing:
        raise SystemExit(f"{SHEET_L1}: no description for {missing}")
    if section_l1 not in l1_order:
        raise SystemExit(f"--section-l1 {section_l1!r} is not an L1 in the workbook")
    names = {s["name"] for s in specialty}
    for value in (fallback, unknown):
        if value not in names:
            raise SystemExit(f"specialty {value!r} is not in the {SHEET_SPECIALTY} sheet")

    # Where an L2 has named L3 children, the fallback is a valid L3 and a blank
    # is not. Where it has none, the blank is the answer.
    for parent, children in l3.items():
        if not any(x["name"] == fallback for x in children):
            l2_entry = next(x for x in l2[parent.split(" > ")[0]] if x["name"] == parent.split(" > ")[1])
            children.append({"name": fallback, "description": f"A {parent.split(' > ')[1]} that none of the "
                             "named L3s fits. Signals that the list needs extending.", "specialty": "",
                             "codeability": l2_entry["codeability"], "codeability_tag": l2_entry["codeability_tag"],
                             "notes": "added by build_taxonomy.py"})

    paths = []
    for a in l1_order:
        if not l2[a]:
            paths.append(a)
        for b in l2[a]:
            children = l3.get(f"{a} > {b['name']}")
            if children:
                paths.extend(f"{a} > {b['name']} > {c['name']}" for c in children)
            else:
                paths.append(f"{a} > {b['name']}")

    return {
        "version": "v24",
        "source": xlsx.name,
        "rules": {
            "section_L1": section_l1,
            "section_L1_note": f"Under {section_l1}, L2 is the content section on the page (axis B), "
                               "decided by section tagging. It is confirmed in the tool and not trained.",
            "fallback": fallback,
            "specialty_unknown": unknown,
            "specialty_other": fallback,
            "specialty_note": f"{unknown!r} means no signal on the page; {fallback!r} means a real service "
                              "line the list does not carry yet, the signal to extend the lexicon.",
            "codeability_values": ["codeable", "non_codeable", "discharge", "needs_review"],
            "codeability_default": "needs_review",
        },
        "L1": [{"name": a, "description": l1_desc[a], **({k: v for k, v in l1_extra[a].items()
                                                         if k != "description"} if a in l1_extra else {})}
               for a in l1_order],
        "L2": l2,
        "L3": l3,
        "specialty": specialty,
        "paths": paths,
        "counts": {"L1": len(l1_order), "L2": sum(len(v) for v in l2.values()),
                   "L3": sum(len(v) for v in l3.values()), "specialty": len(specialty), "paths": len(paths)},
    }


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("xlsx", nargs="?", type=Path, default=HERE / "Page_Taxonomy.xlsx")
    p.add_argument("--section-l1", default="Progress Note", help="the L1 whose L2 is the content section")
    p.add_argument("--fallback", default="Others", help="the catch-all value on every axis")
    p.add_argument("--unknown", default="Unknown", help="the no-signal specialty")
    p.add_argument("--check", action="store_true", help="exit 1 if any copy differs from a fresh build")
    args = p.parse_args(argv)

    data = build(args.xlsx.expanduser(), section_l1=args.section_l1, fallback=args.fallback, unknown=args.unknown)
    text = json.dumps(data, indent=2, ensure_ascii=False) + "\n"
    stale = [c for c in COPIES if not c.is_file() or c.read_text(encoding="utf-8") != text]
    if args.check:
        for c in stale:
            print(f"out of date: {c}")
        return 1 if stale else 0
    for copy in stale:
        copy.write_text(text, encoding="utf-8")
    n = data["counts"]
    print(f"{OUT_NAME}: {n['L1']} L1, {n['L2']} L2, {n['L3']} L3, {n['specialty']} specialties, "
          f"{n['paths']} paths; wrote {len(stale)} of {len(COPIES)} copies")
    return 0


if __name__ == "__main__":
    sys.exit(main())
